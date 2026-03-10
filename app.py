from __future__ import annotations

import os
import re
import sqlite3
import uuid
import zipfile
from urllib import error as url_error
from urllib import request as url_request
import json
from io import BytesIO
from importlib import import_module, util as importlib_util
from calendar import Calendar, monthrange
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Any

from flask import (
    Flask,
    flash,
    g,
    jsonify,
    redirect,
    render_template,
    request,
    session,
    send_file,
    send_from_directory,
    url_for,
)
from werkzeug.security import check_password_hash, generate_password_hash
from werkzeug.utils import secure_filename
from openpyxl import Workbook, load_workbook

BASE_DIR = Path(__file__).resolve().parent
ALLOWED_EXTENSIONS = {"pdf", "png", "jpg", "jpeg", "doc", "docx", "hwp", "txt"}

PROJECT_ROOT = BASE_DIR
TEMPLATE_DIR = PROJECT_ROOT / "templates"
STATIC_DIR = PROJECT_ROOT / "static"
DB_PATH = PROJECT_ROOT / "data" / "counsel.db"
UPLOAD_DIR = PROJECT_ROOT / "data" / "uploads"
BACKUP_DIR = PROJECT_ROOT / "data" / "backups"
MAX_BACKUP_FILES = 20
LOCK_TIMEOUT_SECONDS = 30 * 60
CASE_CONCEPT_THEORIES: list[dict[str, str]] = [
    {"value": "integrative", "label": "통합적 관점(기본)", "guide": "인지·정서·행동·관계·환경 요인을 통합적으로 사례개념화하라."},
    {"value": "cbt", "label": "인지행동치료(CBT)", "guide": "자동사고-핵심신념-행동 패턴의 연쇄와 유지기제를 중심으로 분석하라."},
    {"value": "solution_focused", "label": "해결중심(SFBT)", "guide": "문제보다 예외, 강점, 목표, 작은 변화의 실행계획을 강조하라."},
    {"value": "person_centered", "label": "인간중심", "guide": "공감·수용·진정성 관점에서 정서경험과 관계적 의미를 중심으로 기술하라."},
    {"value": "reality_therapy", "label": "현실치료/선택이론", "guide": "욕구충족, 현재 선택, 책임, 실행가능한 계획(WDEP)을 중심으로 정리하라."},
    {"value": "family_system", "label": "가족체계", "guide": "가족 상호작용, 경계, 의사소통, 반복패턴 등 체계 관점의 유지요인을 포함하라."},
]



def validate_project_layout() -> None:
    required_files = [
        PROJECT_ROOT / "app.py",
        TEMPLATE_DIR / "index.html",
        TEMPLATE_DIR / "base.html",
        STATIC_DIR / "style.css",
    ]
    missing = [str(path) for path in required_files if not path.exists()]
    if missing:
        raise FileNotFoundError(
            "프로젝트 폴더 구조를 찾을 수 없습니다. app.py와 templates/static 파일이 같은 프로젝트 루트에 있어야 합니다. "
            f"누락 파일: {missing}"
        )


def create_app() -> Flask:
    validate_project_layout()

    app = Flask(
        __name__,
        template_folder=str(TEMPLATE_DIR),
        static_folder=str(STATIC_DIR),
    )
    app.secret_key = os.environ.get("SECRET_KEY", "dev-secret-key")

    DB_PATH.parent.mkdir(parents=True, exist_ok=True)
    UPLOAD_DIR.mkdir(parents=True, exist_ok=True)
    BACKUP_DIR.mkdir(parents=True, exist_ok=True)
    init_db()

    @app.before_request
    def before_request() -> None:
        g.db = get_db()
        security_configured = is_security_configured()
        g.security_configured = security_configured

        endpoint = request.endpoint or ""
        exempt_endpoints = {
            "security_setup",
            "unlock_screen",
            "unlock_submit",
            "lock_now",
            "password_reset",
            "static",
        }

        if not security_configured and endpoint not in {"security_setup", "static"}:
            return redirect(url_for("security_setup"))

        if security_configured and endpoint not in exempt_endpoints:
            now_ts = int(datetime.now().timestamp())
            last_activity = int(session.get("last_activity", now_ts))
            is_unlocked = bool(session.get("is_unlocked", False))

            if is_unlocked and now_ts - last_activity > LOCK_TIMEOUT_SECONDS:
                session["is_unlocked"] = False
                flash("30분 이상 활동이 없어 화면이 잠겼습니다.")
                return redirect(url_for("unlock_screen"))

            if not session.get("is_unlocked", False):
                return redirect(url_for("unlock_screen"))

            session["last_activity"] = now_ts

    @app.teardown_request
    def teardown_request(exception: Exception | None) -> None:
        db = g.pop("db", None)
        if db is not None:
            db.close()

    @app.context_processor
    def inject_global_branding() -> dict[str, str]:
        school_name = get_app_setting("school_name", "정동고등학교")
        app_title = f"{school_name} 상담일지 관리"
        return {
            "school_name_global": school_name,
            "app_title": app_title,
        }

    @app.route("/")
    def index() -> str:
        students = query_db(
            "SELECT id, student_no, name, class_name FROM students ORDER BY class_name, name"
        )
        students_count = query_db("SELECT COUNT(*) AS count FROM students", one=True)["count"]
        logs_count = query_db("SELECT COUNT(*) AS count FROM counsel_logs", one=True)["count"]
        latest_logs = query_db(
            """
            SELECT l.id, l.date, l.type, s.name AS student_name, l.summary
            FROM counsel_logs l
            JOIN students s ON s.id = l.student_id
            ORDER BY l.date DESC, l.created_at DESC
            LIMIT 5
            """
        )
        upcoming_schedules = query_db(
            """
            SELECT c.id, c.schedule_date, c.schedule_time, c.title, c.status, s.name AS student_name
            FROM counsel_schedules c
            LEFT JOIN students s ON s.id = c.student_id
            WHERE c.status = 'planned' AND c.schedule_date >= ?
            ORDER BY c.schedule_date ASC, c.schedule_time ASC, c.id ASC
            LIMIT 7
            """,
            (datetime.now().date().isoformat(),),
        )
        backup_files = list_backup_files()
        absence_summary = query_db(
            """
            SELECT
              SUM(CASE WHEN is_active = 1 THEN 1 ELSE 0 END) AS active_count,
              SUM(CASE WHEN is_active = 1 AND home_visit_done = 0 THEN 1 ELSE 0 END) AS home_visit_pending
            FROM unexcused_absences
            """,
            one=True,
        )
        return render_template(
            "index.html",
            students=students,
            students_count=students_count,
            logs_count=logs_count,
            latest_logs=latest_logs,
            upcoming_schedules=upcoming_schedules,
            backup_files=backup_files,
            upcoming_count=len(upcoming_schedules),
            backup_count=len(backup_files),
            active_absence_count=(absence_summary["active_count"] or 0),
            home_visit_pending_count=(absence_summary["home_visit_pending"] or 0),
        )

    @app.route("/backup/create", methods=["POST"])
    def create_backup() -> str:
        backup_name = create_backup_archive(trigger="manual")
        flash(f"백업 파일을 생성했습니다: {backup_name}")
        return redirect(url_for("index"))

    @app.route("/backup/download/<path:filename>")
    def download_backup(filename: str):
        return send_from_directory(BACKUP_DIR, filename, as_attachment=True)

    @app.route("/students/<int:student_id>/ai/case-conceptualization", methods=["POST"])
    def ai_case_conceptualization(student_id: int):
        student = query_db(
            "SELECT id, student_no, name, class_name FROM students WHERE id = ?",
            (student_id,),
            one=True,
        )
        if student is None:
            return jsonify({"ok": False, "error": "학생을 찾을 수 없습니다."}), 404

        payload = request.get_json(silent=True) or {}
        selected_theory = str(payload.get("theory", "integrative")).strip().lower()
        theory_guide = get_case_concept_theory_guide(selected_theory)

        provider = get_ai_provider()
        if provider == "gemini":
            api_key = get_configured_gemini_api_key()
            if not api_key:
                return jsonify({"ok": False, "error": "활성화된 Gemini API 키가 없습니다. AI 설정에서 키를 등록하세요."}), 503
        else:
            api_key = get_configured_openai_api_key()
            if not api_key:
                return jsonify({"ok": False, "error": "활성화된 OpenAI API 키가 없습니다. AI 설정에서 키를 등록하세요."}), 503

        log_rows = query_db(
            """
            SELECT date, type, summary, detail, action_plan, next_date
            FROM counsel_logs
            WHERE student_id = ?
            ORDER BY date ASC, created_at ASC
            """,
            (student_id,),
        )
        if not log_rows:
            return jsonify({"ok": False, "error": "이 학생의 상담일지가 없어 사례개념화를 생성할 수 없습니다."}), 400

        student_label = f"{student['class_name'] or '-'} {student['name']}({student['student_no']})"
        case_context = build_case_context_from_logs(student_label, log_rows)

        system_prompt = (
            "당신은 한국 고등학교 상담교사를 돕는 상담 수퍼바이저다. "
            "제공된 학생의 누적 상담기록을 종합하여 사례개념화 초안을 한국어로 작성하라. "
            f"이론적 기반은 다음 지침을 우선 적용하라: {theory_guide} "
            "출력은 반드시 다음 항목 순서를 지켜라: "
            "1) 핵심문제 2) 경과요약(시간흐름) 3) 유지요인(개인/가정/학교/또래) "
            "4) 보호요인 5) 개입가설 6) 다음회기 질문(3개) 7) 단기개입계획(1~2주)."
        )

        try:
            if provider == "gemini":
                result_text = request_gemini_case_conceptualization(
                    api_key=api_key,
                    system_prompt=system_prompt,
                    user_prompt=case_context,
                )
            else:
                result_text = request_openai_case_conceptualization(
                    api_key=api_key,
                    system_prompt=system_prompt,
                    user_prompt=case_context,
                )
            return jsonify({"ok": True, "result": result_text})
        except RuntimeError as exc:
            return jsonify({"ok": False, "error": str(exc)}), 502


    @app.route("/settings")
    def settings_index() -> str:
        return render_template("settings_index.html")

    @app.route("/settings/ai", methods=["GET", "POST"])
    def ai_settings() -> str:
        if request.method == "POST":
            action = request.form.get("action", "").strip()

            if action == "set_provider":
                provider = request.form.get("provider", "openai").strip().lower()
                if provider not in {"openai", "gemini"}:
                    flash("지원하지 않는 AI 제공자입니다.")
                else:
                    set_app_setting("ai_provider", provider)
                    flash(f"기본 AI 제공자를 {'OpenAI' if provider == 'openai' else 'Gemini'}로 설정했습니다.")
                return redirect(url_for("ai_settings"))

            if action == "clear_openai":
                execute_db("DELETE FROM app_settings WHERE key = ?", ("openai_api_key",))
                flash("OpenAI API 키를 삭제했습니다.")
                return redirect(url_for("ai_settings"))

            if action == "save_openai":
                api_key = request.form.get("openai_api_key", "").strip()
                if not api_key:
                    flash("OpenAI API 키를 입력해주세요.")
                else:
                    set_app_setting("openai_api_key", api_key)
                    flash("OpenAI API 키를 저장했습니다.")
                return redirect(url_for("ai_settings"))

            if action == "clear_gemini":
                execute_db("DELETE FROM app_settings WHERE key = ?", ("gemini_api_key",))
                flash("Gemini API 키를 삭제했습니다.")
                return redirect(url_for("ai_settings"))

            if action == "save_gemini":
                api_key = request.form.get("gemini_api_key", "").strip()
                if not api_key:
                    flash("Gemini API 키를 입력해주세요.")
                else:
                    set_app_setting("gemini_api_key", api_key)
                    flash("Gemini API 키를 저장했습니다.")
                return redirect(url_for("ai_settings"))

        openai_stored_key = get_app_setting("openai_api_key", "").strip()
        gemini_stored_key = get_app_setting("gemini_api_key", "").strip()
        openai_key_masked = f"{'*' * max(len(openai_stored_key) - 4, 0)}{openai_stored_key[-4:]}" if openai_stored_key else ""
        gemini_key_masked = f"{'*' * max(len(gemini_stored_key) - 4, 0)}{gemini_stored_key[-4:]}" if gemini_stored_key else ""

        has_openai_env_key = bool(os.environ.get("OPENAI_API_KEY", "").strip())
        has_gemini_env_key = bool(os.environ.get("GEMINI_API_KEY", "").strip())
        provider = get_ai_provider()

        return render_template(
            "settings_ai.html",
            ai_provider=provider,
            has_openai_stored_key=bool(openai_stored_key),
            openai_key_masked=openai_key_masked,
            has_openai_env_key=has_openai_env_key,
            has_gemini_stored_key=bool(gemini_stored_key),
            gemini_key_masked=gemini_key_masked,
            has_gemini_env_key=has_gemini_env_key,
        )

    @app.route("/settings/ai/validate", methods=["POST"])
    def ai_settings_validate():
        payload = request.get_json(silent=True) or {}
        requested_provider = str(payload.get("provider", "")).strip().lower()
        provider = requested_provider if requested_provider in {"openai", "gemini"} else get_ai_provider()

        if provider == "gemini":
            api_key = get_configured_gemini_api_key()
            if not api_key:
                return jsonify({"ok": False, "message": "활성화된 Gemini API 키가 없습니다."}), 200
            try:
                request_gemini_api_health(api_key)
                return jsonify({"ok": True, "message": "정상: API 키로 Gemini 연결에 성공했습니다."}), 200
            except RuntimeError as exc:
                return jsonify({"ok": False, "message": str(exc)}), 200

        api_key = get_configured_openai_api_key()
        if not api_key:
            return jsonify({"ok": False, "message": "활성화된 OpenAI API 키가 없습니다."}), 200
        try:
            request_openai_api_health(api_key)
            return jsonify({"ok": True, "message": "정상: API 키로 OpenAI 연결에 성공했습니다."}), 200
        except RuntimeError as exc:
            return jsonify({"ok": False, "message": str(exc)}), 200

    @app.route("/setup/security", methods=["GET", "POST"])
    def security_setup() -> str:
        if is_security_configured():
            return redirect(url_for("index"))

        if request.method == "POST":
            password = request.form.get("password", "")
            password_confirm = request.form.get("password_confirm", "")
            reset_phone = normalize_phone(request.form.get("reset_phone", ""))

            if len(password) < 4:
                flash("암호는 4자리 이상으로 설정해 주세요.")
            elif password != password_confirm:
                flash("암호 확인이 일치하지 않습니다.")
            elif len(reset_phone) < 8:
                flash("히든 넘버(전화번호)를 올바르게 입력해 주세요.")
            else:
                set_app_setting("screen_lock_password_hash", generate_password_hash(password))
                set_app_setting("screen_lock_reset_phone", reset_phone)
                session["is_unlocked"] = True
                session["last_activity"] = int(datetime.now().timestamp())
                flash("보안 설정이 완료되었습니다.")
                return redirect(url_for("index"))

        return render_template("security_setup.html")

    @app.route("/lock", methods=["GET"])
    def unlock_screen() -> str:
        if not is_security_configured():
            return redirect(url_for("security_setup"))
        return render_template("lock_screen.html")

    @app.route("/lock", methods=["POST"])
    def unlock_submit() -> str:
        if not is_security_configured():
            return redirect(url_for("security_setup"))

        password = request.form.get("password", "")
        stored_hash = get_app_setting("screen_lock_password_hash", "")
        if not stored_hash or not check_password_hash(stored_hash, password):
            flash("암호가 올바르지 않습니다.")
            return redirect(url_for("unlock_screen"))

        session["is_unlocked"] = True
        session["last_activity"] = int(datetime.now().timestamp())
        flash("잠금이 해제되었습니다.")
        return redirect(url_for("index"))

    @app.route("/lock/now", methods=["POST"])
    def lock_now() -> str:
        session["is_unlocked"] = False
        session["last_activity"] = int(datetime.now().timestamp())
        flash("화면을 잠갔습니다.")
        return redirect(url_for("unlock_screen"))

    @app.route("/lock/reset", methods=["GET", "POST"])
    def password_reset() -> str:
        if not is_security_configured():
            return redirect(url_for("security_setup"))

        if request.method == "POST":
            reset_phone = normalize_phone(request.form.get("reset_phone", ""))
            password = request.form.get("password", "")
            password_confirm = request.form.get("password_confirm", "")

            if len(password) < 4:
                flash("새 암호는 4자리 이상으로 입력해 주세요.")
            elif password != password_confirm:
                flash("새 암호 확인이 일치하지 않습니다.")
            elif reset_phone != get_app_setting("screen_lock_reset_phone", ""):
                flash("히든 넘버(전화번호)가 일치하지 않습니다.")
            else:
                set_app_setting("screen_lock_password_hash", generate_password_hash(password))
                session["is_unlocked"] = True
                session["last_activity"] = int(datetime.now().timestamp())
                flash("암호를 재설정했습니다.")
                return redirect(url_for("index"))

        return render_template("password_reset.html")


    @app.route("/settings/school", methods=["GET", "POST"])
    def school_settings() -> str:
        if request.method == "POST":
            school_name = request.form.get("school_name", "").strip()
            counselor_name = request.form.get("counselor_name", "").strip()

            execute_db(
                """
                INSERT INTO app_settings(key, value)
                VALUES (?, ?)
                ON CONFLICT(key) DO UPDATE SET value = excluded.value
                """,
                ("school_name", school_name),
            )
            execute_db(
                """
                INSERT INTO app_settings(key, value)
                VALUES (?, ?)
                ON CONFLICT(key) DO UPDATE SET value = excluded.value
                """,
                ("counselor_name", counselor_name),
            )
            flash("학교 설정을 저장했습니다.")
            return redirect(url_for("school_settings"))

        school_name = get_app_setting("school_name", "정동고등학교")
        counselor_name = get_app_setting("counselor_name", "전문상담교사")
        return render_template("settings_school.html", school_name=school_name, counselor_name=counselor_name)

    @app.route("/settings/counsel-types", methods=["GET", "POST"])
    def counsel_types_settings() -> str:
        if request.method == "POST":
            action = request.form.get("action", "").strip()
            if action == "add":
                name = request.form.get("name", "").strip()
                if not name:
                    flash("추가할 상담유형 이름을 입력해 주세요.")
                else:
                    try:
                        execute_db("INSERT INTO counsel_types(name, sort_order) VALUES(?, ?)", (name, get_next_counsel_type_order()))
                        flash("상담유형을 추가했습니다.")
                    except sqlite3.IntegrityError:
                        flash("이미 존재하는 상담유형입니다.")
            elif action == "update":
                type_id = request.form.get("type_id", "").strip()
                name = request.form.get("name", "").strip()
                if type_id.isdigit() and name:
                    try:
                        execute_db("UPDATE counsel_types SET name = ? WHERE id = ?", (name, int(type_id)))
                        flash("상담유형을 수정했습니다.")
                    except sqlite3.IntegrityError:
                        flash("이미 존재하는 상담유형입니다.")
                else:
                    flash("수정할 상담유형 정보를 확인해 주세요.")
            elif action == "delete":
                type_id = request.form.get("type_id", "").strip()
                if type_id.isdigit():
                    row = query_db("SELECT name FROM counsel_types WHERE id = ?", (int(type_id),), one=True)
                    if row is not None:
                        used = query_db("SELECT COUNT(*) AS cnt FROM counsel_logs WHERE type = ?", (row["name"],), one=True)
                        if used and used["cnt"] > 0:
                            flash("이미 사용된 상담유형은 삭제할 수 없습니다.")
                        else:
                            execute_db("DELETE FROM counsel_types WHERE id = ?", (int(type_id),))
                            flash("상담유형을 삭제했습니다.")
                else:
                    flash("삭제할 상담유형을 찾을 수 없습니다.")
            return redirect(url_for("counsel_types_settings"))

        counsel_types = get_counsel_types()
        return render_template("counsel_types.html", counsel_types=counsel_types)


    @app.route("/stats")
    def stats() -> str:
        selected_year = request.args.get("year", str(datetime.now().year)).strip()
        if not selected_year.isdigit():
            selected_year = str(datetime.now().year)

        monthly_rows = query_db(
            """
            SELECT substr(date, 1, 7) AS month, COUNT(*) AS count
            FROM counsel_logs
            WHERE substr(date, 1, 4) = ?
            GROUP BY substr(date, 1, 7)
            ORDER BY month
            """,
            (selected_year,),
        )
        monthly_map = {row["month"]: row["count"] for row in monthly_rows}
        monthly_stats = []
        for month in range(1, 13):
            key = f"{selected_year}-{month:02d}"
            monthly_stats.append({"month": key, "count": monthly_map.get(key, 0)})

        yearly_stats = query_db(
            """
            SELECT substr(date, 1, 4) AS year, COUNT(*) AS count
            FROM counsel_logs
            GROUP BY substr(date, 1, 4)
            ORDER BY year DESC
            """
        )

        total_count = query_db("SELECT COUNT(*) AS count FROM counsel_logs", one=True)["count"]
        selected_year_count = sum(row["count"] for row in monthly_stats)

        return render_template(
            "stats.html",
            selected_year=selected_year,
            monthly_stats=monthly_stats,
            yearly_stats=yearly_stats,
            total_count=total_count,
            selected_year_count=selected_year_count,
        )

    @app.route("/stats/neis-monthly-export")
    def stats_neis_monthly_export():
        selected_year = request.args.get("year", str(datetime.now().year)).strip()
        selected_month = request.args.get("month", "").strip()

        if not selected_year.isdigit():
            selected_year = str(datetime.now().year)
        if not selected_month.isdigit() or not (1 <= int(selected_month) <= 12):
            flash("월을 선택해 주세요.")
            return redirect(url_for("stats", year=selected_year))

        month_int = int(selected_month)
        month_key = f"{selected_year}-{month_int:02d}"

        rows = query_db(
            """
            SELECT l.id, l.date, l.type, l.summary, l.detail, l.duration_minutes,
                   l.main_category, l.sub_category, l.media_type,
                   s.student_no, s.grade
            FROM counsel_logs l
            JOIN students s ON s.id = l.student_id
            WHERE substr(l.date, 1, 7) = ?
            ORDER BY l.date ASC, l.created_at ASC
            """,
            (month_key,),
        )

        wb = Workbook()
        ws = wb.active
        ws.title = f"{selected_year}-{month_int:02d}"

        headers = [
            "*상담분류", "*Wee클래스", "*대분류", "*중분류", "*상담구분", "*상담인원", "*학년도", "*상담일자",
            "학년", "성별", "*상담제목", "*상담내용", "*상담시간(시)", "*상담시간(분)", "*상담자소속", "*상담매체구분",
        ]
        ws.append(headers)

        for row in rows:
            duration = int(row["duration_minutes"] or 0)
            hour = duration // 60
            minute = duration % 60
            ws.append([
                "전문상담",          # 고정
                "Wee클래스",         # 고정
                row["main_category"] or "상담",
                row["sub_category"] or "개인상담",
                row["type"] or "",
                1,
                selected_year,
                (row["date"] or "").replace("-", ""),
                row["grade"] or "",
                "",
                row["summary"] or "",
                row["summary"] or "",
                hour,
                minute,
                "전문상담교사",
                row["media_type"] or "면담",
            ])

        for col in ws.columns:
            max_len = 0
            col_letter = col[0].column_letter
            for cell in col:
                value = "" if cell.value is None else str(cell.value)
                max_len = max(max_len, len(value))
            ws.column_dimensions[col_letter].width = min(max_len + 2, 40)

        out = BytesIO()
        wb.save(out)
        out.seek(0)

        filename = f"neis_monthly_counsel_{selected_year}_{month_int:02d}.xlsx"
        return send_file(
            out,
            mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            as_attachment=True,
            download_name=filename,
        )

    @app.route("/stats/annual-ledger")
    def stats_annual_ledger() -> str:
        selected_year = request.args.get("year", str(datetime.now().year)).strip()
        if not selected_year.isdigit():
            selected_year = str(datetime.now().year)

        rows = query_db(
            """
            SELECT l.id, l.date, l.type, l.summary, l.detail, l.action_plan,
                   l.counsel_period, l.duration_minutes,
                   s.name AS student_name, s.grade, s.class_no
            FROM counsel_logs l
            JOIN students s ON s.id = l.student_id
            WHERE substr(l.date, 1, 4) = ?
            ORDER BY l.date ASC, l.created_at ASC
            """,
            (selected_year,),
        )

        return render_template(
            "stats_annual_ledger_print.html",
            selected_year=selected_year,
            rows=rows,
            printed_at=datetime.now().strftime("%Y-%m-%d %H:%M"),
        )


    @app.route("/schedule", methods=["GET", "POST"])
    def schedule() -> str:
        requested_date = request.args.get("date", "").strip()
        try:
            selected_date_obj = date.fromisoformat(requested_date) if requested_date else datetime.now().date()
        except ValueError:
            selected_date_obj = datetime.now().date()
        selected_date = selected_date_obj.isoformat()

        month_param = request.args.get("month", "").strip()
        try:
            if month_param:
                year, month = month_param.split("-")
                calendar_year = int(year)
                calendar_month = int(month)
            else:
                calendar_year = selected_date_obj.year
                calendar_month = selected_date_obj.month
        except ValueError:
            calendar_year = selected_date_obj.year
            calendar_month = selected_date_obj.month

        first_day_of_month = date(calendar_year, calendar_month, 1)
        last_day_of_month = date(calendar_year, calendar_month, monthrange(calendar_year, calendar_month)[1])
        prev_month = date(calendar_year - 1, 12, 1) if calendar_month == 1 else date(calendar_year, calendar_month - 1, 1)
        next_month = date(calendar_year + 1, 1, 1) if calendar_month == 12 else date(calendar_year, calendar_month + 1, 1)

        student_rows = query_db("SELECT id, student_no, name, grade, class_no, class_name FROM students ORDER BY grade, class_no, name")
        grades, grade_class_map = get_grade_class_filters(student_rows)
        counsel_types = get_counsel_types()
        edit_schedule_id = request.values.get("edit_id", "").strip()
        edit_schedule = None
        if edit_schedule_id.isdigit():
            edit_schedule = query_db(
                """
                SELECT id, student_id, schedule_date, schedule_time, title, note, status
                FROM counsel_schedules
                WHERE id = ?
                """,
                (int(edit_schedule_id),),
                one=True,
            )

        if request.method == "POST":
            student_id = request.form.get("student_id", "").strip()
            schedule_date = request.form.get("schedule_date", "").strip()
            schedule_slot = request.form.get("schedule_slot", "").strip()
            custom_slot = request.form.get("custom_slot", "").strip()
            title = request.form.get("title", "").strip()
            note = request.form.get("note", "").strip()

            schedule_time = schedule_slot
            if schedule_slot == "직접입력":
                schedule_time = custom_slot

            if not schedule_date or not title:
                flash("상담일과 일정 제목은 필수입니다.")
            elif schedule_slot == "직접입력" and not custom_slot:
                flash("직접입력을 선택한 경우 시간/교시를 입력해 주세요.")
            else:
                if edit_schedule and edit_schedule["status"] == "planned":
                    execute_db(
                        """
                        UPDATE counsel_schedules
                        SET student_id = ?, schedule_date = ?, schedule_time = ?, title = ?, note = ?
                        WHERE id = ?
                        """,
                        (
                            student_id or None,
                            schedule_date,
                            schedule_time or None,
                            title,
                            note,
                            edit_schedule["id"],
                        ),
                    )
                    flash("상담 일정이 수정되었습니다.")
                else:
                    execute_db(
                        """
                        INSERT INTO counsel_schedules(student_id, schedule_date, schedule_time, title, note, status, created_at)
                        VALUES (?, ?, ?, ?, ?, 'planned', ?)
                        """,
                        (
                            student_id or None,
                            schedule_date,
                            schedule_time or None,
                            title,
                            note,
                            datetime.now().isoformat(timespec="seconds"),
                        ),
                    )
                    flash("상담 일정이 등록되었습니다.")
                return redirect(url_for("schedule", date=schedule_date, month=schedule_date[:7]))

        schedules_for_date = query_db(
            """
            SELECT c.id, c.schedule_date, c.schedule_time, c.title, c.note, c.status,
                   s.id AS student_id, s.name AS student_name, s.grade, s.class_no, s.class_name
            FROM counsel_schedules c
            LEFT JOIN students s ON s.id = c.student_id
            WHERE c.schedule_date = ?
            ORDER BY c.schedule_time ASC, c.id ASC
            """,
            (selected_date,),
        )

        month_rows = query_db(
            """
            SELECT c.id, c.schedule_date, c.schedule_time, c.title, c.status,
                   s.name AS student_name
            FROM counsel_schedules c
            LEFT JOIN students s ON s.id = c.student_id
            WHERE c.schedule_date BETWEEN ? AND ?
            ORDER BY c.schedule_date ASC, c.schedule_time ASC, c.id ASC
            """,
            (first_day_of_month.isoformat(), last_day_of_month.isoformat()),
        )

        month_counts: dict[str, int] = {}
        schedules_by_date: dict[str, list[dict[str, Any]]] = {}
        for row in month_rows:
            day_key = row["schedule_date"]
            month_counts[day_key] = month_counts.get(day_key, 0) + 1
            schedules_by_date.setdefault(day_key, []).append(dict(row))

        cal = Calendar(firstweekday=0)
        calendar_dates = [d for week in cal.monthdatescalendar(calendar_year, calendar_month) for d in week]
        holiday_years = {d.year for d in calendar_dates}
        auto_holidays = get_korean_public_holidays(holiday_years)
        temp_holiday_rows = query_db("SELECT holiday_date, name FROM school_holidays")
        holiday_map: dict[str, str] = dict(auto_holidays)
        for row in temp_holiday_rows:
            holiday_map[row["holiday_date"]] = row["name"] or "임시공휴일"

        calendar_weeks: list[list[dict[str, Any]]] = []
        for week_dates in cal.monthdatescalendar(calendar_year, calendar_month):
            week_items: list[dict[str, Any]] = []
            for day_obj in week_dates:
                day_key = day_obj.isoformat()
                day_schedules = schedules_by_date.get(day_key, [])
                week_items.append(
                    {
                        "date": day_key,
                        "day": day_obj.day,
                        "is_current_month": day_obj.month == calendar_month,
                        "is_selected": day_key == selected_date,
                        "is_today": day_key == datetime.now().date().isoformat(),
                        "is_holiday": day_key in holiday_map,
                        "holiday_name": holiday_map.get(day_key, ""),
                        "schedules": day_schedules[:3],
                        "extra_count": max(0, len(day_schedules) - 3),
                    }
                )
            calendar_weeks.append(week_items)

        upcoming_schedules = query_db(
            """
            SELECT c.id, c.schedule_date, c.schedule_time, c.title, c.status,
                   s.id AS student_id, s.name AS student_name, s.grade, s.class_no
            FROM counsel_schedules c
            LEFT JOIN students s ON s.id = c.student_id
            WHERE c.status = 'planned' AND c.schedule_date >= ?
            ORDER BY c.schedule_date ASC, c.schedule_time ASC, c.id ASC
            LIMIT 20
            """,
            (datetime.now().date().isoformat(),),
        )

        return render_template(
            "schedule.html",
            students=student_rows,
            selected_date=selected_date,
            selected_month=first_day_of_month.strftime("%Y-%m"),
            month_title=first_day_of_month.strftime("%Y년 %m월"),
            prev_month=prev_month.strftime("%Y-%m"),
            next_month=next_month.strftime("%Y-%m"),
            calendar_weeks=calendar_weeks,
            schedules_for_date=schedules_for_date,
            month_counts=month_counts,
            upcoming_schedules=upcoming_schedules,
            edit_schedule=edit_schedule,
            grades=grades,
            grade_class_map=grade_class_map,
        )

    @app.route("/absence", methods=["GET", "POST"])
    def absence_tracker() -> str:
        if request.method == "POST":
            action = request.form.get("action", "").strip()
            if action == "add_absence":
                student_id = request.form.get("student_id", "").strip()
                start_date = request.form.get("start_date", "").strip()
                if not student_id or not student_id.isdigit() or not start_date:
                    flash("학생과 시작일을 입력해 주세요.")
                else:
                    try:
                        date.fromisoformat(start_date)
                        execute_db(
                            """
                            INSERT INTO unexcused_absences(student_id, start_date, home_visit_done, created_at)
                            VALUES (?, ?, 0, ?)
                            """,
                            (int(student_id), start_date, datetime.now().isoformat(timespec="seconds")),
                        )
                        flash("미인정결석 학생을 등록했습니다.")
                        return redirect(url_for("absence_tracker"))
                    except ValueError:
                        flash("시작일 형식이 올바르지 않습니다.")
            elif action == "add_holiday":
                holiday_date = request.form.get("holiday_date", "").strip()
                holiday_name = request.form.get("holiday_name", "").strip()
                if not holiday_date:
                    flash("임시공휴일 날짜를 입력해 주세요.")
                else:
                    try:
                        date.fromisoformat(holiday_date)
                        execute_db(
                            """
                            INSERT INTO school_holidays(holiday_date, name)
                            VALUES (?, ?)
                            ON CONFLICT(holiday_date) DO UPDATE SET name = excluded.name
                            """,
                            (holiday_date, holiday_name or "공휴일"),
                        )
                        flash("임시공휴일을 저장했습니다.")
                        return redirect(url_for("absence_tracker"))
                    except ValueError:
                        flash("임시공휴일 날짜 형식이 올바르지 않습니다.")
            elif action == "delete_holiday":
                holiday_date = request.form.get("holiday_date", "").strip()
                if holiday_date:
                    execute_db("DELETE FROM school_holidays WHERE holiday_date = ?", (holiday_date,))
                    flash("임시공휴일을 삭제했습니다.")
                    return redirect(url_for("absence_tracker"))
            elif action == "mark_return":
                absence_id = request.form.get("absence_id", "").strip()
                if absence_id.isdigit():
                    execute_db(
                        "UPDATE unexcused_absences SET is_active = 0, return_date = ? WHERE id = ?",
                        (datetime.now().date().isoformat(), int(absence_id)),
                    )
                    flash("복귀 처리되었습니다.")
                    return redirect(url_for("absence_tracker"))
            elif action == "update_absence":
                absence_id = request.form.get("absence_id", "").strip()
                start_date = request.form.get("start_date", "").strip()
                if not absence_id.isdigit() or not start_date:
                    flash("수정할 대상과 시작일을 확인해 주세요.")
                else:
                    try:
                        date.fromisoformat(start_date)
                        execute_db(
                            "UPDATE unexcused_absences SET start_date = ? WHERE id = ?",
                            (start_date, int(absence_id)),
                        )
                        flash("시작일을 수정했습니다.")
                        return redirect(url_for("absence_tracker"))
                    except ValueError:
                        flash("시작일 형식이 올바르지 않습니다.")
            elif action == "delete_absence":
                absence_id = request.form.get("absence_id", "").strip()
                if absence_id.isdigit():
                    execute_db("DELETE FROM unexcused_absences WHERE id = ?", (int(absence_id),))
                    flash("대상자를 삭제했습니다.")
                    return redirect(url_for("absence_tracker"))

        student_rows = query_db(
            "SELECT id, student_no, name, grade, class_no, class_name FROM students ORDER BY grade, class_no, name"
        )
        temp_holiday_rows = query_db("SELECT holiday_date, name FROM school_holidays ORDER BY holiday_date")

        years = {datetime.now().year}
        for r in query_db("SELECT start_date FROM unexcused_absences"):
            try:
                years.add(date.fromisoformat(r["start_date"]).year)
            except ValueError:
                continue
        auto_holidays = get_korean_public_holidays(years)
        temp_holidays = {row["holiday_date"]: (row["name"] or "임시공휴일") for row in temp_holiday_rows}
        holidays = set(auto_holidays.keys()) | set(temp_holidays.keys())

        rows = query_db(
            """
            SELECT a.id, a.student_id, a.start_date, a.home_visit_done, a.home_visit_done_at, a.is_active, a.return_date, a.created_at,
                   s.student_no, s.name AS student_name, s.grade, s.class_no, s.class_name
            FROM unexcused_absences a
            JOIN students s ON s.id = a.student_id
            ORDER BY a.is_active DESC, a.start_date ASC, a.id DESC
            """
        )

        today = datetime.now().date()
        tracked = []
        for row in rows:
            try:
                start = date.fromisoformat(row["start_date"])
            except ValueError:
                continue
            absence_days = business_days_count(start, today, holidays)
            danger_ratio = min(max(absence_days / 7.0, 0), 1)
            tracked.append({
                "row": row,
                "absence_days": absence_days,
                "is_report_due": absence_days >= 7,
                "is_home_visit_due": absence_days >= 2,
                "danger_ratio": danger_ratio,
                "is_active": bool(row["is_active"] if row["is_active"] is not None else 1),
            })

        return render_template(
            "absence_tracker.html",
            students=student_rows,
            tracked=tracked,
            temp_holiday_rows=temp_holiday_rows,
            auto_holiday_count=len(auto_holidays),
            today=today.isoformat(),
        )

    @app.route("/absence/<int:absence_id>/home-visit", methods=["POST"])
    def mark_home_visit(absence_id: int) -> str:
        row = query_db("SELECT id, home_visit_done, is_active FROM unexcused_absences WHERE id = ?", (absence_id,), one=True)
        if row is None:
            flash("대상을 찾을 수 없습니다.")
            return redirect(url_for("absence_tracker"))

        if not row["is_active"]:
            flash("이미 복귀 처리된 대상입니다.")
            return redirect(url_for("absence_tracker"))

        if row["home_visit_done"]:
            flash("이미 가정방문 완료로 처리되었습니다.")
            return redirect(url_for("absence_tracker"))

        execute_db(
            "UPDATE unexcused_absences SET home_visit_done = 1, home_visit_done_at = ? WHERE id = ?",
            (datetime.now().isoformat(timespec="seconds"), absence_id),
        )
        flash("가정방문 완료 처리되었습니다.")
        return redirect(url_for("absence_tracker"))

    @app.route("/students", methods=["GET", "POST"])
    def students() -> str:
        if request.method == "POST":
            form_action = request.form.get("form_action", "manual").strip()
            if form_action == "import_excel":
                excel_file = request.files.get("student_excel")
                if not excel_file or excel_file.filename == "":
                    flash("엑셀 파일을 선택해 주세요.")
                else:
                    try:
                        imported, skipped = import_students_from_excel(excel_file.stream)
                        flash(f"엑셀 일괄등록 완료: {imported}건 반영, {skipped}건 건너뜀")
                        return redirect(url_for("students"))
                    except Exception as exc:
                        flash(f"엑셀 처리 중 오류가 발생했습니다: {exc}")
            else:
                student_no = request.form.get("student_no", "").strip()
                name = request.form.get("name", "").strip()
                grade = request.form.get("grade", "").strip()
                class_no = request.form.get("class_no", "").strip()
                class_name = request.form.get("class_name", "").strip() or (f"{grade}-{class_no}" if grade and class_no else "")
                homeroom_teacher = request.form.get("homeroom_teacher", "").strip()
                phone = request.form.get("phone", "").strip()
                guardian_phone = request.form.get("guardian_phone", "").strip()
                note = request.form.get("note", "").strip()

                if not student_no or not name:
                    flash("학번과 이름은 필수입니다.")
                else:
                    try:
                        execute_db(
                            """
                            INSERT INTO students(student_no, name, grade, class_no, class_name, homeroom_teacher, phone, guardian_phone, note)
                            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
                            """,
                            (student_no, name, grade or None, class_no or None, class_name or None, homeroom_teacher or None, phone, guardian_phone, note),
                        )
                        flash("학생 정보가 저장되었습니다.")
                        return redirect(url_for("students"))
                    except sqlite3.IntegrityError:
                        flash("이미 등록된 학번입니다.")

        q = request.args.get("q", "").strip()
        if q:
            student_rows = query_db(
                """
                SELECT id, student_no, name, grade, class_no, class_name, homeroom_teacher, phone, guardian_phone, note
                FROM students
                WHERE student_no LIKE ? OR name LIKE ? OR class_name LIKE ? OR grade LIKE ? OR class_no LIKE ? OR homeroom_teacher LIKE ?
                ORDER BY grade, class_no, name
                """,
                (f"%{q}%", f"%{q}%", f"%{q}%", f"%{q}%", f"%{q}%", f"%{q}%"),
            )
        else:
            student_rows = query_db(
                """
                SELECT id, student_no, name, grade, class_no, class_name, homeroom_teacher, phone, guardian_phone, note
                FROM students
                ORDER BY grade, class_no, name
                """
            )
        grades, grade_class_map = get_grade_class_filters(student_rows)
        return render_template("students.html", students=student_rows, q=q, grades=grades, grade_class_map=grade_class_map)

    @app.route("/logs")
    def logs_list() -> str:
        q = request.args.get("q", "").strip()
        if q:
            rows = query_db(
                """
                SELECT l.id, l.date, l.type, l.summary, s.id AS student_id, s.name AS student_name,
                       s.grade, s.class_no, s.class_name
                FROM counsel_logs l
                JOIN students s ON s.id = l.student_id
                WHERE l.summary LIKE ? OR l.type LIKE ? OR s.name LIKE ? OR l.date LIKE ?
                ORDER BY l.date DESC, l.created_at DESC
                """,
                (f"%{q}%", f"%{q}%", f"%{q}%", f"%{q}%"),
            )
        else:
            rows = query_db(
                """
                SELECT l.id, l.date, l.type, l.summary, s.id AS student_id, s.name AS student_name,
                       s.grade, s.class_no, s.class_name
                FROM counsel_logs l
                JOIN students s ON s.id = l.student_id
                ORDER BY l.date DESC, l.created_at DESC
                """
            )
        return render_template("logs.html", logs=rows, q=q)

    @app.route("/logs/new", methods=["GET", "POST"])
    def new_log() -> str:
        student_rows = query_db("SELECT id, student_no, name, grade, class_no, class_name FROM students ORDER BY grade, class_no, name")
        grades, grade_class_map = get_grade_class_filters(student_rows)
        schedule_id = request.values.get("schedule_id", "").strip()
        counsel_types = get_counsel_types()

        prefill_schedule = None
        if schedule_id.isdigit():
            prefill_schedule = query_db(
                """
                SELECT c.id, c.schedule_date, c.schedule_time, c.title, c.note, c.status,
                       s.id AS student_id, s.name AS student_name, s.grade, s.class_no, s.class_name
                FROM counsel_schedules c
                LEFT JOIN students s ON s.id = c.student_id
                WHERE c.id = ?
                """,
                (int(schedule_id),),
                one=True,
            )

        if request.method == "POST":
            student_id = request.form.get("student_id", "").strip()
            date = request.form.get("date", "").strip()
            log_type = request.form.get("type", "").strip()
            summary = request.form.get("summary", "").strip()
            detail = request.form.get("detail", "").strip()
            action_plan = request.form.get("action_plan", "").strip()
            main_category = request.form.get("main_category", "").strip() or "상담"
            sub_category = request.form.get("sub_category", "").strip() or "개인상담"
            media_type = request.form.get("media_type", "").strip() or "면담"
            counsel_period = request.form.get("counsel_period", "").strip()
            duration_minutes_raw = request.form.get("duration_minutes", "").strip()
            next_date = request.form.get("next_date", "").strip()
            linked_schedule_id = request.form.get("schedule_id", "").strip()

            duration_minutes = None
            if duration_minutes_raw:
                try:
                    duration_minutes = int(duration_minutes_raw)
                    if duration_minutes <= 0:
                        raise ValueError
                except ValueError:
                    flash("소요시간은 1 이상의 숫자(분)로 입력해 주세요.")
                    return redirect(url_for("new_log", schedule_id=linked_schedule_id))

            if not student_id or not date or not log_type or not summary:
                flash("학생, 상담일, 상담유형, 상담제목은 필수입니다.")
            else:
                log_id = execute_db(
                    """
                    INSERT INTO counsel_logs(
                        student_id, date, type, summary, detail, action_plan,
                        main_category, sub_category, media_type,
                        counsel_period, duration_minutes, next_date, created_at
                    )
                    VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                    """,
                    (
                        student_id,
                        date,
                        log_type,
                        summary,
                        detail,
                        action_plan,
                        main_category,
                        sub_category,
                        media_type,
                        counsel_period or None,
                        duration_minutes,
                        next_date or None,
                        datetime.now().isoformat(timespec="seconds"),
                    ),
                )

                files = request.files.getlist("attachments")
                for file in files:
                    if not file or file.filename == "":
                        continue
                    if allowed_file(file.filename):
                        original_name = file.filename
                        ext = original_name.rsplit(".", 1)[1].lower()
                        new_name = f"{uuid.uuid4().hex}.{ext}"
                        save_path = UPLOAD_DIR / secure_filename(new_name)
                        file.save(save_path)
                        execute_db(
                            """
                            INSERT INTO attachments(log_id, file_path, original_name, created_at)
                            VALUES (?, ?, ?, ?)
                            """,
                            (
                                log_id,
                                save_path.name,
                                secure_filename(original_name),
                                datetime.now().isoformat(timespec="seconds"),
                            ),
                        )

                if linked_schedule_id.isdigit():
                    execute_db(
                        """
                        UPDATE counsel_schedules
                        SET status = 'done'
                        WHERE id = ?
                        """,
                        (int(linked_schedule_id),),
                    )
                    flash("연결된 상담 일정이 완료 처리되었습니다.")

                if next_date:
                    execute_db(
                        """
                        INSERT INTO counsel_schedules(student_id, schedule_date, schedule_time, title, note, status, created_at)
                        VALUES (?, ?, ?, ?, ?, 'planned', ?)
                        """,
                        (
                            student_id,
                            next_date,
                            None,
                            f"후속상담 - {summary}",
                            action_plan,
                            datetime.now().isoformat(timespec="seconds"),
                        ),
                    )
                    flash("다음 상담일이 일정에 자동 등록되었습니다.")

                flash("상담일지가 저장되었습니다.")
                return redirect(url_for("student_detail", student_id=student_id))

        return render_template(
            "new_log.html",
            students=student_rows,
            prefill_schedule=prefill_schedule,
            grades=grades,
            grade_class_map=grade_class_map,
            counsel_types=counsel_types,
        )

    @app.route("/logs/<int:log_id>/confirmation")
    def log_confirmation(log_id: int) -> str:
        row = query_db(
            """
            SELECT l.id, l.date, l.type, l.summary, l.detail, l.counsel_period,
                   s.id AS student_id, s.name AS student_name, s.grade, s.class_no, s.class_name
            FROM counsel_logs l
            JOIN students s ON s.id = l.student_id
            WHERE l.id = ?
            """,
            (log_id,),
            one=True,
        )
        if row is None:
            flash("상담일지를 찾을 수 없습니다.")
            return redirect(url_for("logs_list"))

        return render_template(
            "log_confirmation_print.html",
            log=row,
            printed_at=datetime.now().strftime("%Y-%m-%d"),
            school_name=get_app_setting("school_name", "정동고등학교"),
            counselor_name=get_app_setting("counselor_name", "전문상담교사"),
        )

    @app.route("/logs/<int:log_id>/edit", methods=["GET", "POST"])
    def edit_log(log_id: int) -> str:
        log_row = query_db(
            """
            SELECT id, student_id, date, type, summary, detail, action_plan, main_category, sub_category, media_type, counsel_period, duration_minutes, next_date
            FROM counsel_logs
            WHERE id = ?
            """,
            (log_id,),
            one=True,
        )
        if log_row is None:
            flash("상담일지를 찾을 수 없습니다.")
            return redirect(url_for("students"))

        student_rows = query_db(
            "SELECT id, student_no, name, grade, class_no, class_name FROM students ORDER BY grade, class_no, name"
        )
        counsel_types = get_counsel_types()
        grades, grade_class_map = get_grade_class_filters(student_rows)

        if request.method == "POST":
            student_id = request.form.get("student_id", "").strip()
            log_date = request.form.get("date", "").strip()
            log_type = request.form.get("type", "").strip()
            summary = request.form.get("summary", "").strip()
            detail = request.form.get("detail", "").strip()
            action_plan = request.form.get("action_plan", "").strip()
            main_category = request.form.get("main_category", "").strip() or "상담"
            sub_category = request.form.get("sub_category", "").strip() or "개인상담"
            media_type = request.form.get("media_type", "").strip() or "면담"
            counsel_period = request.form.get("counsel_period", "").strip()
            duration_minutes_raw = request.form.get("duration_minutes", "").strip()
            next_date = request.form.get("next_date", "").strip()

            duration_minutes = None
            if duration_minutes_raw:
                try:
                    duration_minutes = int(duration_minutes_raw)
                    if duration_minutes <= 0:
                        raise ValueError
                except ValueError:
                    flash("소요시간은 1 이상의 숫자(분)로 입력해 주세요.")
                    return redirect(url_for("edit_log", log_id=log_id))

            if not student_id or not log_date or not log_type or not summary:
                flash("학생, 상담일, 상담유형, 상담제목은 필수입니다.")
            else:
                execute_db(
                    """
                    UPDATE counsel_logs
                    SET student_id = ?, date = ?, type = ?, summary = ?, detail = ?, action_plan = ?,
                        main_category = ?, sub_category = ?, media_type = ?,
                        counsel_period = ?, duration_minutes = ?, next_date = ?
                    WHERE id = ?
                    """,
                    (
                        student_id,
                        log_date,
                        log_type,
                        summary,
                        detail,
                        action_plan,
                        main_category,
                        sub_category,
                        media_type,
                        counsel_period or None,
                        duration_minutes,
                        next_date or None,
                        log_id,
                    ),
                )
                flash("상담일지를 수정했습니다.")
                return redirect(url_for("student_detail", student_id=student_id))

        attachments = query_db(
            "SELECT id, file_path, original_name FROM attachments WHERE log_id = ? ORDER BY id",
            (log_id,),
        )
        return render_template(
            "edit_log.html",
            log=log_row,
            attachments=attachments,
            students=student_rows,
            grades=grades,
            grade_class_map=grade_class_map,
            counsel_types=counsel_types,
        )

    @app.route("/logs/<int:log_id>/delete", methods=["POST"])
    def delete_log(log_id: int) -> str:
        log_row = query_db("SELECT id, student_id FROM counsel_logs WHERE id = ?", (log_id,), one=True)
        if log_row is None:
            flash("상담일지를 찾을 수 없습니다.")
            return redirect(url_for("students"))

        files = query_db("SELECT file_path FROM attachments WHERE log_id = ?", (log_id,))
        for file_row in files:
            (UPLOAD_DIR / file_row["file_path"]).unlink(missing_ok=True)

        execute_db("DELETE FROM attachments WHERE log_id = ?", (log_id,))
        execute_db("DELETE FROM counsel_logs WHERE id = ?", (log_id,))
        flash("상담일지를 삭제했습니다.")
        return redirect(url_for("student_detail", student_id=log_row["student_id"]))


    @app.route("/students/<int:student_id>/update", methods=["POST"])
    def update_student(student_id: int) -> str:
        student = query_db("SELECT id FROM students WHERE id = ?", (student_id,), one=True)
        if student is None:
            flash("학생을 찾을 수 없습니다.")
            return redirect(url_for("students"))

        student_no = request.form.get("student_no", "").strip()
        phone = request.form.get("phone", "").strip()
        guardian_phone = request.form.get("guardian_phone", "").strip()
        note = request.form.get("note", "").strip()

        if not student_no:
            flash("학번은 비워둘 수 없습니다.")
            return redirect(url_for("student_detail", student_id=student_id))

        try:
            execute_db(
                """
                UPDATE students
                SET student_no = ?, phone = ?, guardian_phone = ?, note = ?
                WHERE id = ?
                """,
                (student_no, phone, guardian_phone, note, student_id),
            )
            flash("학생 정보를 수정했습니다.")
        except sqlite3.IntegrityError:
            flash("이미 등록된 학번입니다.")

        return redirect(url_for("student_detail", student_id=student_id))

    @app.route("/students/<int:student_id>/record-form")
    def student_record_form(student_id: int) -> str:
        student = query_db(
            "SELECT id, student_no, name, grade, class_no, class_name, homeroom_teacher, phone, guardian_phone, note FROM students WHERE id = ?",
            (student_id,),
            one=True,
        )
        if student is None:
            flash("학생을 찾을 수 없습니다.")
            return redirect(url_for("students"))

        logs = query_db(
            """
            SELECT id, date, type, summary, detail, action_plan,
                   main_category, sub_category, media_type,
                   counsel_period, duration_minutes, next_date
            FROM counsel_logs
            WHERE student_id = ?
            ORDER BY date ASC, created_at ASC
            """,
            (student_id,),
        )

        counseling_dates = [row["date"] for row in logs[:18]]
        date_slots_first = counseling_dates[:9]
        date_slots_second = counseling_dates[9:18]
        while len(date_slots_first) < 9:
            date_slots_first.append("")
        while len(date_slots_second) < 9:
            date_slots_second.append("")

        counselor_name = get_app_setting("counselor_name", student["homeroom_teacher"] or "-")
        school_name = get_app_setting("school_name", "정동고등학교")
        end_date = logs[-1]["date"] if logs else ""
        closing_comment = logs[-1]["action_plan"] if logs and logs[-1]["action_plan"] else ""

        return render_template(
            "student_record_form_print.html",
            student=student,
            logs=logs,
            date_slots_first=date_slots_first,
            date_slots_second=date_slots_second,
            counselor_name=counselor_name,
            school_name=school_name,
            end_date=end_date,
            closing_comment=closing_comment,
            printed_at=datetime.now().strftime("%Y-%m-%d %H:%M"),
        )


    @app.route("/students/<int:student_id>/print")
    def student_print(student_id: int) -> str:
        student = query_db(
            "SELECT id, student_no, name, grade, class_no, class_name, homeroom_teacher, phone, guardian_phone, note FROM students WHERE id = ?",
            (student_id,),
            one=True,
        )
        if student is None:
            flash("학생을 찾을 수 없습니다.")
            return redirect(url_for("students"))

        date_from = request.args.get("date_from", "").strip()
        date_to = request.args.get("date_to", "").strip()

        query = (
            """
            SELECT id, date, type, summary, detail, action_plan, main_category, sub_category, media_type, counsel_period, duration_minutes, next_date, created_at
            FROM counsel_logs
            WHERE student_id = ?
            """
        )
        params: list[Any] = [student_id]
        if date_from:
            query += " AND date >= ?"
            params.append(date_from)
        if date_to:
            query += " AND date <= ?"
            params.append(date_to)

        query += " ORDER BY date DESC, created_at DESC"
        logs = query_db(query, tuple(params))

        return render_template(
            "student_print.html",
            student=student,
            logs=logs,
            date_from=date_from,
            date_to=date_to,
            printed_at=datetime.now().strftime("%Y-%m-%d %H:%M"),
        )


    @app.route("/students/<int:student_id>/case-concept", methods=["POST"])
    def save_case_concept(student_id: int) -> str:
        student = query_db("SELECT id FROM students WHERE id = ?", (student_id,), one=True)
        if student is None:
            flash("학생을 찾을 수 없습니다.")
            return redirect(url_for("students"))

        content = request.form.get("case_concept", "").strip()
        if not content:
            flash("사례개념화 내용을 입력해 주세요.")
            return redirect(url_for("student_detail", student_id=student_id))

        execute_db(
            """
            INSERT INTO student_case_concepts(student_id, content, updated_at)
            VALUES (?, ?, ?)
            ON CONFLICT(student_id) DO UPDATE SET
                content = excluded.content,
                updated_at = excluded.updated_at
            """,
            (student_id, content, datetime.now().isoformat(timespec="seconds")),
        )
        flash("사례개념화를 저장했습니다.")
        return redirect(url_for("student_detail", student_id=student_id))

    @app.route("/students/<int:student_id>")
    def student_detail(student_id: int) -> str:
        student = query_db(
            "SELECT id, student_no, name, grade, class_no, class_name, homeroom_teacher, phone, guardian_phone, note FROM students WHERE id = ?",
            (student_id,),
            one=True,
        )
        if student is None:
            flash("학생을 찾을 수 없습니다.")
            return redirect(url_for("students"))

        log_rows = query_db(
            """
            SELECT id, date, type, summary, detail, action_plan, main_category, sub_category, media_type, counsel_period, duration_minutes, next_date, created_at
            FROM counsel_logs
            WHERE student_id = ?
            ORDER BY date DESC, created_at DESC
            """,
            (student_id,),
        )
        concept_row = query_db(
            "SELECT content, updated_at FROM student_case_concepts WHERE student_id = ?",
            (student_id,),
            one=True,
        )

        logs_with_files: list[dict[str, Any]] = []
        for row in log_rows:
            files = query_db(
                "SELECT id, file_path, original_name FROM attachments WHERE log_id = ? ORDER BY id",
                (row["id"],),
            )
            logs_with_files.append({"log": row, "files": files})

        return render_template(
            "student_detail.html",
            student=student,
            logs=logs_with_files,
            case_concept=concept_row,
            case_concept_theories=CASE_CONCEPT_THEORIES,
        )

    @app.route("/uploads/<path:filename>")
    def uploaded_file(filename: str):
        return send_from_directory(UPLOAD_DIR, filename, as_attachment=True)

    @app.route("/init-db")
    def init_db_route() -> str:
        init_db()
        flash("데이터베이스를 초기화했습니다.")
        return redirect(url_for("index"))

    return app


def get_db() -> sqlite3.Connection:
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    return conn


def init_db() -> None:
    conn = get_db()
    with conn:
        conn.executescript(
            """
            CREATE TABLE IF NOT EXISTS students (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                student_no TEXT UNIQUE NOT NULL,
                name TEXT NOT NULL,
                grade TEXT,
                class_no TEXT,
                class_name TEXT,
                homeroom_teacher TEXT,
                phone TEXT,
                guardian_phone TEXT,
                note TEXT
            );

            CREATE TABLE IF NOT EXISTS counsel_logs (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                student_id INTEGER NOT NULL,
                date TEXT NOT NULL,
                type TEXT NOT NULL,
                summary TEXT NOT NULL,
                detail TEXT,
                action_plan TEXT,
                main_category TEXT,
                sub_category TEXT,
                media_type TEXT,
                counsel_period TEXT,
                duration_minutes INTEGER,
                next_date TEXT,
                created_at TEXT NOT NULL,
                FOREIGN KEY(student_id) REFERENCES students(id)
            );

            CREATE TABLE IF NOT EXISTS attachments (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                log_id INTEGER NOT NULL,
                file_path TEXT NOT NULL,
                original_name TEXT NOT NULL,
                created_at TEXT NOT NULL,
                FOREIGN KEY(log_id) REFERENCES counsel_logs(id)
            );

            CREATE TABLE IF NOT EXISTS counsel_schedules (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                student_id INTEGER,
                schedule_date TEXT NOT NULL,
                schedule_time TEXT,
                title TEXT NOT NULL,
                note TEXT,
                status TEXT NOT NULL DEFAULT 'planned',
                created_at TEXT NOT NULL,
                FOREIGN KEY(student_id) REFERENCES students(id)
            );

            CREATE TABLE IF NOT EXISTS app_settings (
                key TEXT PRIMARY KEY,
                value TEXT NOT NULL
            );

            CREATE TABLE IF NOT EXISTS student_case_concepts (
                student_id INTEGER PRIMARY KEY,
                content TEXT NOT NULL,
                updated_at TEXT NOT NULL,
                FOREIGN KEY(student_id) REFERENCES students(id)
            );

            CREATE TABLE IF NOT EXISTS unexcused_absences (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                student_id INTEGER NOT NULL,
                start_date TEXT NOT NULL,
                home_visit_done INTEGER NOT NULL DEFAULT 0,
                home_visit_done_at TEXT,
                is_active INTEGER NOT NULL DEFAULT 1,
                return_date TEXT,
                created_at TEXT NOT NULL,
                FOREIGN KEY(student_id) REFERENCES students(id)
            );

            CREATE TABLE IF NOT EXISTS school_holidays (
                holiday_date TEXT PRIMARY KEY,
                name TEXT
            );

            CREATE TABLE IF NOT EXISTS counsel_types (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                name TEXT UNIQUE NOT NULL,
                sort_order INTEGER NOT NULL DEFAULT 0
            );
            """
        )
        add_column_if_missing(conn, "students", "grade", "TEXT")
        add_column_if_missing(conn, "students", "class_no", "TEXT")
        add_column_if_missing(conn, "students", "homeroom_teacher", "TEXT")
        add_column_if_missing(conn, "counsel_logs", "main_category", "TEXT")
        add_column_if_missing(conn, "counsel_logs", "sub_category", "TEXT")
        add_column_if_missing(conn, "counsel_logs", "media_type", "TEXT")
        add_column_if_missing(conn, "counsel_logs", "counsel_period", "TEXT")
        add_column_if_missing(conn, "counsel_logs", "duration_minutes", "INTEGER")
        add_column_if_missing(conn, "unexcused_absences", "is_active", "INTEGER NOT NULL DEFAULT 1")
        add_column_if_missing(conn, "unexcused_absences", "return_date", "TEXT")
        seed_default_counsel_types(conn)
    conn.close()



def add_column_if_missing(conn: sqlite3.Connection, table: str, column: str, col_type: str) -> None:
    rows = conn.execute(f"PRAGMA table_info({table})").fetchall()
    existing = {row[1] for row in rows}
    if column not in existing:
        conn.execute(f"ALTER TABLE {table} ADD COLUMN {column} {col_type}")


def get_korean_public_holidays(years: set[int]) -> dict[str, str]:
    years = {y for y in years if isinstance(y, int)}
    if not years:
        years = {datetime.now().year}

    # `holidays` 패키지가 설치된 경우 한국 공휴일(대체공휴일 포함)을 사용
    if importlib_util.find_spec("holidays") is not None:
        holidays_mod = import_module("holidays")
        kr = holidays_mod.KR(years=sorted(years))
        return {d.isoformat(): str(name) for d, name in kr.items()}

    # 패키지가 없을 때 최소한의 고정 공휴일 fallback
    fixed = {}
    for y in years:
        for md, name in [
            ("01-01", "신정"),
            ("03-01", "삼일절"),
            ("05-05", "어린이날"),
            ("06-06", "현충일"),
            ("08-15", "광복절"),
            ("10-03", "개천절"),
            ("10-09", "한글날"),
            ("12-25", "성탄절"),
        ]:
            fixed[f"{y}-{md}"] = name
    return fixed


def business_days_count(start: date, end: date, holiday_dates: set[str]) -> int:
    if end < start:
        return 0
    count = 0
    current = start
    while current <= end:
        iso = current.isoformat()
        if current.weekday() < 5 and iso not in holiday_dates:
            count += 1
        current += timedelta(days=1)
    return count


def parse_grade_class_from_text(text: str) -> tuple[str, str]:
    match = re.search(r"(\d+)\s*[-반]\s*(\d+)", text)
    if not match:
        return "", ""
    return match.group(1), match.group(2)


def parse_homeroom_teacher(ws) -> str:
    for row in ws.iter_rows(min_row=1, max_row=8, min_col=1, max_col=8, values_only=True):
        for value in row:
            if not value:
                continue
            text = str(value).strip()
            match = re.search(r"담임\s*[:：]?\s*([^\s]+)", text)
            if match:
                return match.group(1).strip()
    return ""


def import_students_from_excel(file_stream) -> tuple[int, int]:
    wb = load_workbook(file_stream, data_only=True)
    inserted_or_updated = 0
    skipped = 0

    for ws in wb.worksheets:
        grade, class_no = parse_grade_class_from_text(ws.title)
        teacher = parse_homeroom_teacher(ws)
        class_name = f"{grade}-{class_no}" if grade and class_no else ws.title

        for row in ws.iter_rows(min_row=1, max_col=3, values_only=True):
            raw_no = row[0]
            raw_name = row[1]
            if raw_no is None or raw_name is None:
                continue

            no_text = str(raw_no).strip()
            name = str(raw_name).strip()
            if not re.fullmatch(r"\d+", no_text):
                continue
            if not re.fullmatch(r"[가-힣A-Za-z]{2,20}", name):
                continue

            no = int(no_text)
            student_no = f"{grade or 'X'}{int(class_no) if class_no.isdigit() else 0:02d}{no:02d}"

            execute_db(
                """
                INSERT INTO students(student_no, name, grade, class_no, class_name, homeroom_teacher)
                VALUES (?, ?, ?, ?, ?, ?)
                ON CONFLICT(student_no) DO UPDATE SET
                    name = excluded.name,
                    grade = excluded.grade,
                    class_no = excluded.class_no,
                    class_name = excluded.class_name,
                    homeroom_teacher = excluded.homeroom_teacher
                """,
                (student_no, name, grade or None, class_no or None, class_name, teacher or None),
            )
            inserted_or_updated += 1

    return inserted_or_updated, skipped


def get_grade_class_filters(student_rows) -> tuple[list[str], dict[str, list[str]]]:
    grades = sorted({(row["grade"] or "").strip() for row in student_rows if row["grade"]})
    mapping: dict[str, set[str]] = {g: set() for g in grades}
    for row in student_rows:
        grade = (row["grade"] or "").strip()
        class_no = (row["class_no"] or "").strip()
        if grade and class_no:
            mapping.setdefault(grade, set()).add(class_no)
    class_map = {k: sorted(v, key=lambda x: int(x) if x.isdigit() else x) for k, v in mapping.items()}
    return grades, class_map



def seed_default_counsel_types(conn: sqlite3.Connection) -> None:
    defaults = [
        "학업",
        "진로",
        "성격",
        "성",
        "대인관계",
        "가정 및 가족관계",
        "일탈 및 비행",
        "학교폭력 가해",
        "학교폭력 피해",
        "자해 및 자살",
        "정신건강",
        "컴퓨터 및 스마트폰 과사용",
        "정보제공",
        "기타",
    ]
    legacy_only = {"진로", "학업", "정서", "생활지도", "기타"}

    existing_rows = conn.execute("SELECT id, name FROM counsel_types ORDER BY sort_order, id").fetchall()
    existing_names = [row[1] for row in existing_rows]

    if not existing_rows:
        for idx, name in enumerate(defaults, start=1):
            conn.execute("INSERT INTO counsel_types(name, sort_order) VALUES(?, ?)", (name, idx))
        return

    # 테이블이 사실상 예전 기본유형만 있다면 전체를 새 기본유형으로 교체
    if set(existing_names).issubset(legacy_only):
        conn.execute("DELETE FROM counsel_types")
        for idx, name in enumerate(defaults, start=1):
            conn.execute("INSERT INTO counsel_types(name, sort_order) VALUES(?, ?)", (name, idx))
        return

    # 새 기본유형은 항상 존재/정렬되도록 보장(사용자 정의 유형은 유지)
    for idx, name in enumerate(defaults, start=1):
        row = conn.execute("SELECT id FROM counsel_types WHERE name = ?", (name,)).fetchone()
        if row:
            conn.execute("UPDATE counsel_types SET sort_order = ? WHERE id = ?", (idx, row[0]))
        else:
            conn.execute("INSERT INTO counsel_types(name, sort_order) VALUES(?, ?)", (name, idx))

    # 예전 기본유형 중 현재 기본목록에 없는 값은 미사용일 때만 정리
    for legacy_name in ["정서", "생활지도"]:
        used = conn.execute("SELECT COUNT(*) FROM counsel_logs WHERE type = ?", (legacy_name,)).fetchone()[0]
        if used == 0:
            conn.execute("DELETE FROM counsel_types WHERE name = ?", (legacy_name,))


def get_counsel_types() -> list[sqlite3.Row]:
    return query_db("SELECT id, name, sort_order FROM counsel_types ORDER BY sort_order, id")


def get_next_counsel_type_order() -> int:
    row = query_db("SELECT COALESCE(MAX(sort_order), 0) + 1 AS next_order FROM counsel_types", one=True)
    return row["next_order"] if row else 1


def list_backup_files() -> list[Path]:
    BACKUP_DIR.mkdir(parents=True, exist_ok=True)
    return sorted(BACKUP_DIR.glob("counsel_backup_*.zip"), reverse=True)


def prune_old_backups() -> None:
    backups = list_backup_files()
    for old_file in backups[MAX_BACKUP_FILES:]:
        old_file.unlink(missing_ok=True)


def create_backup_archive(trigger: str = "auto") -> str:
    BACKUP_DIR.mkdir(parents=True, exist_ok=True)
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    backup_name = f"counsel_backup_{timestamp}_{trigger}.zip"
    backup_path = BACKUP_DIR / backup_name

    with zipfile.ZipFile(backup_path, "w", compression=zipfile.ZIP_DEFLATED) as zip_file:
        if DB_PATH.exists():
            zip_file.write(DB_PATH, arcname="counsel.db")
        if UPLOAD_DIR.exists():
            for file_path in UPLOAD_DIR.rglob("*"):
                if file_path.is_file():
                    arcname = Path("uploads") / file_path.relative_to(UPLOAD_DIR)
                    zip_file.write(file_path, arcname=str(arcname))

    prune_old_backups()
    return backup_name


def ensure_daily_auto_backup() -> str | None:
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    conn.execute(
        """
        CREATE TABLE IF NOT EXISTS app_settings (
            key TEXT PRIMARY KEY,
            value TEXT NOT NULL
        )
        """
    )
    today = datetime.now().date().isoformat()
    row = conn.execute("SELECT value FROM app_settings WHERE key = 'last_auto_backup_date'").fetchone()
    if row and row["value"] == today:
        conn.close()
        return None

    backup_name = create_backup_archive(trigger="auto")
    conn.execute(
        """
        INSERT INTO app_settings(key, value)
        VALUES ('last_auto_backup_date', ?)
        ON CONFLICT(key) DO UPDATE SET value = excluded.value
        """,
        (today,),
    )
    conn.commit()
    conn.close()
    return backup_name


def anonymize_text_for_ai(text: str) -> str:
    masked = text
    masked = re.sub(r"\b\d{4,}\b", "[숫자정보]", masked)
    masked = re.sub(r"\d{2,3}[-\s]?\d{3,4}[-\s]?\d{4}", "[연락처]", masked)
    masked = re.sub(r"[가-힣]{2,4}(?=\s?(학생|군|양|님))", "[학생]", masked)
    return masked




def build_case_context_from_logs(student_label: str, log_rows: list[sqlite3.Row]) -> str:
    lines = [f"[학생정보] {anonymize_text_for_ai(student_label)}", "[누적 상담일지]"]
    for idx, row in enumerate(log_rows, start=1):
        lines.append(
            '\n'.join(
                [
                    f"- 기록 {idx}",
                    f"  일자: {row['date']}",
                    f"  유형: {row['type']}",
                    f"  요약: {anonymize_text_for_ai(row['summary'] or '')}",
                    f"  상세: {anonymize_text_for_ai(row['detail'] or '')}",
                    f"  조치/계획: {anonymize_text_for_ai(row['action_plan'] or '')}",
                    f"  다음상담일: {row['next_date'] or '-'}",
                ]
            )
        )
    return '\n'.join(lines)


def request_openai_api_health(api_key: str) -> None:
    req = url_request.Request(
        "https://api.openai.com/v1/models",
        method="GET",
        headers={"Authorization": f"Bearer {api_key}"},
    )
    try:
        with url_request.urlopen(req, timeout=20):
            return
    except url_error.HTTPError as exc:
        detail = exc.read().decode("utf-8", errors="ignore")
        raise RuntimeError(f"OpenAI API 키 검증 실패: {detail}") from exc
    except url_error.URLError as exc:
        raise RuntimeError("OpenAI API 연결 실패: 네트워크를 확인하세요.") from exc


def get_app_setting(key: str, default: str = "") -> str:
    row = query_db("SELECT value FROM app_settings WHERE key = ?", (key,), one=True)
    if row and row["value"] is not None:
        return str(row["value"])
    return default


def set_app_setting(key: str, value: str) -> None:
    execute_db(
        """
        INSERT INTO app_settings(key, value)
        VALUES (?, ?)
        ON CONFLICT(key) DO UPDATE SET value = excluded.value
        """,
        (key, value),
    )


def normalize_phone(raw: str) -> str:
    return re.sub(r"\D", "", raw or "")


def is_security_configured() -> bool:
    return bool(get_app_setting("screen_lock_password_hash", "").strip())


def get_case_concept_theory_guide(theory_value: str) -> str:
    for theory in CASE_CONCEPT_THEORIES:
        if theory["value"] == theory_value:
            return theory["guide"]
    return CASE_CONCEPT_THEORIES[0]["guide"]


def get_ai_provider() -> str:
    provider = get_app_setting("ai_provider", "openai").strip().lower()
    return provider if provider in {"openai", "gemini"} else "openai"


def get_configured_gemini_api_key() -> str:
    row = query_db("SELECT value FROM app_settings WHERE key = ?", ("gemini_api_key",), one=True)
    if row and row["value"].strip():
        return row["value"].strip()
    return os.environ.get("GEMINI_API_KEY", "").strip()


def get_configured_openai_api_key() -> str:
    row = query_db("SELECT value FROM app_settings WHERE key = ?", ("openai_api_key",), one=True)
    if row and row["value"].strip():
        return row["value"].strip()
    return os.environ.get("OPENAI_API_KEY", "").strip()

def request_openai_case_conceptualization(api_key: str, system_prompt: str, user_prompt: str) -> str:
    model = os.environ.get("OPENAI_MODEL", "gpt-4.1-mini")
    body = {
        "model": model,
        "temperature": 0.4,
        "messages": [
            {"role": "system", "content": system_prompt},
            {"role": "user", "content": user_prompt},
        ],
    }
    data = json.dumps(body).encode("utf-8")
    req = url_request.Request(
        "https://api.openai.com/v1/chat/completions",
        data=data,
        method="POST",
        headers={
            "Content-Type": "application/json",
            "Authorization": f"Bearer {api_key}",
        },
    )
    try:
        with url_request.urlopen(req, timeout=30) as response:
            raw = response.read().decode("utf-8")
            parsed = json.loads(raw)
            return parsed["choices"][0]["message"]["content"].strip()
    except url_error.HTTPError as exc:
        detail = exc.read().decode("utf-8", errors="ignore")
        raise RuntimeError(f"OpenAI API 오류: {detail}") from exc
    except url_error.URLError as exc:
        raise RuntimeError("OpenAI API 연결에 실패했습니다. 네트워크를 확인하세요.") from exc


def request_gemini_api_health(api_key: str) -> None:
    req = url_request.Request(
        f"https://generativelanguage.googleapis.com/v1beta/models?key={api_key}",
        method="GET",
    )
    try:
        with url_request.urlopen(req, timeout=20):
            return
    except url_error.HTTPError as exc:
        detail = exc.read().decode("utf-8", errors="ignore")
        raise RuntimeError(f"Gemini API 키 검증 실패: {detail}") from exc
    except url_error.URLError as exc:
        raise RuntimeError("Gemini API 연결 실패: 네트워크를 확인하세요.") from exc


def request_gemini_case_conceptualization(api_key: str, system_prompt: str, user_prompt: str) -> str:
    configured_model = os.environ.get("GEMINI_MODEL", "").strip()
    candidate_models: list[str] = []
    if configured_model:
        candidate_models.append(configured_model)
    candidate_models.extend([
        "gemini-2.0-flash",
        "gemini-1.5-flash-latest",
        "gemini-1.5-pro-latest",
    ])

    seen: set[str] = set()
    deduped_models: list[str] = []
    for model in candidate_models:
        if model and model not in seen:
            seen.add(model)
            deduped_models.append(model)

    body = {
        "contents": [
            {"role": "user", "parts": [{"text": system_prompt + "\n\n" + user_prompt}]}
        ],
        "generationConfig": {
            "temperature": 0.4,
        },
    }
    data = json.dumps(body).encode("utf-8")

    last_error_detail = ""
    for model in deduped_models:
        req = url_request.Request(
            f"https://generativelanguage.googleapis.com/v1beta/models/{model}:generateContent?key={api_key}",
            data=data,
            method="POST",
            headers={"Content-Type": "application/json"},
        )
        try:
            with url_request.urlopen(req, timeout=30) as response:
                raw = response.read().decode("utf-8")
                parsed = json.loads(raw)
                candidates = parsed.get("candidates") or []
                if not candidates:
                    raise RuntimeError("Gemini 응답에 생성 결과가 없습니다.")
                parts = (candidates[0].get("content") or {}).get("parts") or []
                text = "".join(str(p.get("text", "")) for p in parts).strip()
                if not text:
                    raise RuntimeError("Gemini 응답 텍스트를 파싱하지 못했습니다.")
                return text
        except url_error.HTTPError as exc:
            detail = exc.read().decode("utf-8", errors="ignore")
            last_error_detail = detail
            lowered = detail.lower()
            if exc.code == 404 or "not_found" in lowered or "not found for api version" in lowered:
                continue
            raise RuntimeError(f"Gemini API 오류: {detail}") from exc
        except url_error.URLError as exc:
            raise RuntimeError("Gemini API 연결에 실패했습니다. 네트워크를 확인하세요.") from exc

    hint = "GEMINI_MODEL 환경변수에 사용 가능한 모델명을 지정해 보세요. (예: gemini-2.0-flash)"
    raise RuntimeError(f"Gemini API 오류: 사용 가능한 모델을 찾지 못했습니다. {hint} 상세: {last_error_detail}")


def query_db(query: str, params: tuple[Any, ...] = (), one: bool = False):
    cur = g.db.execute(query, params)
    rows = cur.fetchall()
    cur.close()
    return (rows[0] if rows else None) if one else rows


def execute_db(query: str, params: tuple[Any, ...] = ()) -> int:
    cur = g.db.execute(query, params)
    g.db.commit()
    row_id = cur.lastrowid
    cur.close()
    return row_id


def allowed_file(filename: str) -> bool:
    return "." in filename and filename.rsplit(".", 1)[1].lower() in ALLOWED_EXTENSIONS


app = create_app()

if __name__ == "__main__":
    init_db()
    auto_backup = ensure_daily_auto_backup()
    if auto_backup:
        print(f"[INFO] 자동 백업 생성: {auto_backup}")
    app.run(host="0.0.0.0", port=5000, debug=True)
